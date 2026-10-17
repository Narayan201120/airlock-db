"""Convert the public MoneyData workbook to the built-in JSONL schema for testing.

This preparation script is optional and is not part of the AirLock DB runtime.
It requires openpyxl to read the source XLSX file.
"""

from __future__ import annotations

import argparse
import json
import os
import tempfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from uuid import NAMESPACE_URL, uuid5

try:
    from openpyxl import load_workbook
except ImportError as exc:
    raise SystemExit(
        "MoneyData XLSX preparation requires openpyxl. Install it only in the "
        "data-preparation environment; AirLock DB does not require it."
    ) from exc


_SHEET_NAME = "Anonymized Original with Catego"
_EXPECTED_HEADERS = (
    "Transaction Number",
    "Transaction Date",
    "Transaction Type",
    "Transaction Description",
    "Debit Amount",
    "Credit Amount",
    "Balance",
    "Category",
    "Location City",
    "Location Country",
)
_COMPARE_CHUNK_SIZE = 1024 * 1024


def _transaction_date(value: object, row_number: int) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return datetime.strptime(value, "%d/%m/%Y").date()
        except ValueError as exc:
            raise ValueError(f"row {row_number} has an unsupported transaction date") from exc
    raise ValueError(f"row {row_number} has no transaction date")


def _amount(value: object) -> str | None:
    if value is None:
        return None
    return format(Decimal(str(value)), "f")


def _same_content(left: Path, right: Path) -> bool:
    with left.open("rb") as left_stream, right.open("rb") as right_stream:
        while True:
            left_chunk = left_stream.read(_COMPARE_CHUNK_SIZE)
            right_chunk = right_stream.read(_COMPARE_CHUNK_SIZE)
            if left_chunk != right_chunk:
                return False
            if not left_chunk:
                return True


def convert(workbook_path: Path, output_path: Path) -> dict[str, int]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    counts = {"records": 0, "debit": 0, "credit": 0, "unmapped_amount": 0}
    try:
        sheet = workbook[_SHEET_NAME]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if headers != _EXPECTED_HEADERS:
            raise ValueError(f"unexpected MoneyData headers: {headers!r}")

        output_path.parent.mkdir(parents=True, exist_ok=True)
        descriptor, temporary_name = tempfile.mkstemp(
            prefix=f".{output_path.name}.", suffix=".partial", dir=output_path.parent
        )
        os.close(descriptor)
        temporary_path = Path(temporary_name)
        try:
            with temporary_path.open("w", encoding="utf-8", newline="\n") as output:
                for row_number, values in enumerate(rows, start=2):
                    if not any(value is not None for value in values):
                        continue
                    if len(values) != len(_EXPECTED_HEADERS):
                        raise ValueError(f"row {row_number} has an unexpected column count")

                    (
                        source_number,
                        raw_date,
                        _source_type,
                        raw_description,
                        debit,
                        credit,
                        _balance,
                        category,
                        _city,
                        _country,
                    ) = values

                    transaction_id = (
                        str(
                            uuid5(
                                NAMESPACE_URL,
                                f"mendeley:dnxtg6n4rv:transaction:{source_number}",
                            )
                        )
                        if source_number is not None
                        else None
                    )
                    if debit is not None and credit is None:
                        amount = _amount(debit)
                        transaction_type = "DEBIT"
                        counts["debit"] += 1
                    elif credit is not None and debit is None:
                        amount = _amount(credit)
                        transaction_type = "CREDIT"
                        counts["credit"] += 1
                    else:
                        amount = None
                        transaction_type = "UNKNOWN"
                        counts["unmapped_amount"] += 1

                    transaction_day = _transaction_date(raw_date, row_number)
                    record = {
                        "transaction_id": transaction_id,
                        "account_id": "MONEYDATA-SOURCE",
                        "posted_at": f"{transaction_day.isoformat()}T00:00:00+00:00",
                        "amount": amount,
                        "currency": "GBP",
                        "transaction_type": transaction_type,
                        "status": "POSTED",
                        "counterparty": raw_description,
                        "description": category,
                    }
                    output.write(
                        json.dumps(record, ensure_ascii=False, allow_nan=False) + "\n"
                    )
                    counts["records"] += 1

            if output_path.exists():
                if _same_content(temporary_path, output_path):
                    return counts
                raise FileExistsError(
                    f"refusing to overwrite different existing output: {output_path}"
                )
            try:
                os.link(temporary_path, output_path)
            except FileExistsError:
                if not _same_content(temporary_path, output_path):
                    raise FileExistsError(
                        f"refusing to overwrite different existing output: {output_path}"
                    )
        finally:
            temporary_path.unlink(missing_ok=True)
    finally:
        workbook.close()
    return counts


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Prepare public MoneyData rows for an AirLock DB built-in-schema test."
    )
    parser.add_argument("workbook", type=Path, help="downloaded MoneyData XLSX file")
    parser.add_argument("output", type=Path, help="JSONL output path")
    args = parser.parse_args()
    try:
        counts = convert(args.workbook, args.output)
    except (OSError, ValueError) as exc:
        parser.error(str(exc))
    for name, value in counts.items():
        print(f"{name}={value}")
    print(f"jsonl={args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
